"""Gulf of Cadiz food web (Web of Life FW_005): clean, build, and profile
a weighted directed predator-prey network.

Source: Torres et al. 2013, Ecological Modelling 265:26-44, via the Web of
Life database (www.web-of-life.es). Run from anywhere:

    python code/build_network.py

Reads  data/web-of-life_2026-10-15_210043/FW_005.csv
Writes output/  (culled matrix, cleaning log, node traits, edge list,
                 GraphML, summary, and all of it as one Excel workbook)
       figures/ (diet matrix heatmap, food web drawn by trophic level)
"""
from itertools import combinations
from pathlib import Path
import numpy as np
import pandas as pd
import networkx as nx
import matplotlib.pyplot as plt
from matplotlib.colors import LogNorm, LinearSegmentedColormap, to_rgb
from matplotlib.patches import Rectangle
from matplotlib.transforms import Bbox
from matplotlib.lines import Line2D
from mpl_toolkits.axes_grid1 import make_axes_locatable

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "web-of-life_2026-10-15_210043"
OUT = ROOT / "output"
FIG = ROOT / "figures"
OUT.mkdir(exist_ok=True)
FIG.mkdir(exist_ok=True)

report = []
def log(*a):
    line = " ".join(str(x) for x in a)
    print(line)
    report.append(line)

cleaning = []
def record(action, target, detail, links_before, links_after):
    cleaning.append(dict(step=len(cleaning) + 1, action=action, target=target,
                         detail=detail, links_before=links_before,
                         links_after=links_after))

# ---------------------------------------------------------------- load
# FW_005.csv is a labelled 44x44 matrix. Rows are prey (resources), columns
# are predators (consumers), and each cell is the fraction of the column
# group's diet made up of the row group (Ecopath diet composition).
raw = pd.read_csv(RAW / "FW_005.csv", index_col=0)
assert raw.shape == (44, 44)
assert list(raw.index) == list(raw.columns), "row and column labels differ"
names = list(raw.index)
M = raw.to_numpy(dtype=float)
idx = {n: i for i, n in enumerate(names)}
L0 = int((M > 0).sum())

log("== raw ==")
log(f"groups: {len(names)}   links (nonzero cells): {L0}")

# every consumer's diet should sum to 1, which also confirms columns = predators
colsum = M.sum(axis=0)
consumers = [n for n in names if colsum[idx[n]] > 0]
dev = max(abs(colsum[idx[n]] - 1) for n in consumers)
assert dev < 0.01
log(f"consumers: {len(consumers)}   max |diet sum - 1| = {dev:.4f}")
record("check", "all consumer columns",
       f"{len(consumers)} consumer diets each sum to 1 (max deviation "
       f"{dev:.4f}), so rows = prey and columns = predators. No change.",
       L0, L0)
diet_total_original = dict(zip(names, colsum))

# ---------------------------------------------------------------- clean
# "Import" is not a functional group. In Ecopath it is the share of a
# predator's diet eaten outside the model area. The paper's model has 43
# groups; Import is the 44th row (and an all-zero column: it eats nothing).
log("\n== cleaning ==")
imp = M[idx["Import"]]
import_share = {n: float(imp[idx[n]]) for n in names if imp[idx[n]] > 0}
for n, s in import_share.items():
    log(f"Import in diet of {n}: {s:.3f}")
assert M[:, idx["Import"]].sum() == 0
keep = [n for n in names if n != "Import"]
M = M[np.ix_([idx[n] for n in keep], [idx[n] for n in keep])]
names = keep
idx = {n: i for i, n in enumerate(names)}
L1 = int((M > 0).sum())
record("remove", "Import (row and column)",
       "Not a functional group: Ecopath's share of a predator's diet eaten "
       "outside the model area. Nonzero only as prey of "
       + ", ".join(f"{n} ({s:.3f})" for n, s in import_share.items())
       + ". Its column is all zero.", L0, L1)
log(f"dropped Import -> {len(names)} groups, {L1} links")

record("check", "exports / emigration",
       "Searched for any other row or column that is not a living or "
       "detrital group (e.g. export, emigration, fishery). None: Import was "
       "the only boundary flow in the matrix. Fishery catches and exports "
       "live in the paper's tables, not in the diet matrix.", L1, L1)
record("renormalise", "diets of " + " and ".join(import_share),
       "Without Import these diets sum to "
       + ", ".join(f"{1 - s:.3f}" for s in import_share.values())
       + ". For trophic level, omnivory index and diet diversity every diet is "
       "rescaled to sum to 1 (column diet_share_renorm in the edge list), which "
       "assumes food eaten outside the area has the same make-up as the local "
       "diet. The culled matrix keeps the original shares.", L1, L1)

NONLIVING = {"Detritus", "Discard"}
VERTEBRATES = {"Seabirds", "Killer whales", "Dolphins", "Loggerhead turtles"}
FISHES = {"Sharks", "Demersal piscivores", "Anglerfishes", "Skates",
          "Adult hake", "Juvenile hake", "Mackerels", "Horse mackerels",
          "Deep-sea fishes", "Large benthopelagic fishes",
          "Small benthopelagic fishes", "Flatfishes", "Demersal fishes",
          "Mullets", "Commercial sparids 1", "Commercial sparids 2",
          "Blue whiting", "Small demersal fishes", "Anchovy", "Sardine"}
def group(n):
    if n in NONLIVING:   return "non-living"
    if n in VERTEBRATES: return "marine mammals, birds, turtles"
    if n in FISHES:      return "fishes"
    return "invertebrates & plankton"
assert sum(group(n) == "invertebrates & plankton" for n in names) == 17

# ---------------------------------------------------------------- build
# Edge direction prey -> predator, i.e. the direction energy flows.
# Reverse the graph (G.reverse()) for the "who eats whom" convention.
G = nx.DiGraph()
for n in names:
    G.add_node(n, group=group(n))
for i, j in zip(*np.nonzero(M)):
    G.add_edge(names[i], names[j], weight=float(M[i, j]))

N, L = G.number_of_nodes(), G.number_of_edges()
loops = sorted(nx.nodes_with_selfloops(G))
log("\n== network ==")
log(f"nodes {N}   edges {L}   self-loops (cannibalism) {len(loops)}")
log("cannibals:", ", ".join(loops))
log(f"isolates: {list(nx.isolates(G)) or 'none'}")
log(f"weakly connected components: {nx.number_weakly_connected_components(G)}")
scc = max(nx.strongly_connected_components(G), key=len)
log(f"largest strongly connected component: {len(scc)} of {N}")
log(f"directed connectance L/N^2 = {L / N**2:.3f}   "
    f"(Web of Life's 0.437 uses L/(S(S-1)/2) with S=44)")

w = np.array([d["weight"] for *_, d in G.edges(data=True)])
log(f"weights: min {w.min():g}  median {np.median(w):.3f}  max {w.max():g}")
log(f"links < 0.01: {(w < 0.01).sum()}   < 0.001: {(w < 0.001).sum()}")

record("keep", "Detritus, Discard",
       "Non-living basal groups kept as nodes. Detritus is the most-eaten "
       "group in the web (summed diet share 5.13). Removing both leaves 41 nodes / 396 links, still one "
       "weakly connected component with no isolates.", L, L)
record("keep", f"{len(loops)} cannibalistic self-loops",
       "Kept in the matrix and the graph. Excluded when counting prey and "
       "predators (n_prey, n_predators) and from diet diversity.", L, L)
record("keep", "weak links",
       f"No weight threshold: smallest share {w.min():g}; "
       f"{(w < 0.01).sum()} links below 0.01 and {(w < 0.001).sum()} below "
       "0.001 are kept.", L, L)

# ---------------------------------------------------------------- nodes
# Trophic level, Ecopath style: TL_i = 1 + sum_j D_ji TL_j over prey j, with
# basal groups = 1 and D the diet matrix rescaled so every diet sums to 1.
D = M / np.where(M.sum(axis=0) > 0, M.sum(axis=0), 1)   # columns sum to 1
TL = np.linalg.solve(np.eye(N) - D.T, np.ones(N))

# Omnivory index (Christensen & Pauly 1992): diet-weighted variance of the
# prey trophic levels around TL_i - 1. Zero for a consumer feeding on one level.
OI = np.array([(D[:, i] * (TL - (TL[i] - 1)) ** 2).sum() for i in range(N)])

# Diet diversity: exp(Shannon entropy) of the diet without cannibalism,
# i.e. "how many equally used prey" the diet is equivalent to.
def effective_prey(i):
    p = np.delete(D[:, i], i)
    p = p[p > 0] / p[p > 0].sum() if p.sum() > 0 else p[p > 0]
    return float(np.exp(-(p * np.log(p)).sum())) if len(p) else np.nan

in_loop = set().union(*(c for c in nx.strongly_connected_components(G)
                        if len(c) > 1))
nodes = pd.DataFrame({
    "category": [group(n) for n in names],
    "living": [n not in NONLIVING for n in names],
    "trophic_level": TL.round(3),
    "omnivory_index": OI.round(3),
    # generality / vulnerability, not counting the self-loop
    "n_prey": [G.in_degree(n) - (n in loops) for n in names],
    "n_predators": [G.out_degree(n) - (n in loops) for n in names],
    "effective_prey": [round(effective_prey(i), 3) for i in range(N)],
    # out-strength: summed share of predator diets this group makes up
    "diet_share_sum": [round(G.out_degree(n, weight="weight"), 4) for n in names],
    "cannibal": [n in loops for n in names],
    "cannibalism_share": [float(M[i, i]) for i in range(N)],
    "import_share_removed": [import_share.get(n, 0.0) for n in names],
    "diet_total_original": [round(float(diet_total_original[n]), 4) for n in names],
    "in_feeding_loop": [n in in_loop for n in names],
}, index=pd.Index(names, name="group_name"))
nodes.insert(2, "role", np.select(
    [nodes.n_prey == 0, nodes.n_predators == 0],
    ["basal", "top predator"], "intermediate"))
nodes.insert(7, "total_degree", nodes.n_prey + nodes.n_predators)
nodes.loc[nodes.role == "basal", "effective_prey"] = np.nan

log("\nbasal:", ", ".join(nodes.index[nodes.role == "basal"]))
log("top predators (eaten by no other group):",
    ", ".join(nodes.index[nodes.role == "top predator"]))
log("\nmost prey groups (generalists):")
for n, r in nodes.nlargest(5, "n_prey").iterrows():
    log(f"  {n:28s} {r.n_prey}")
log("most predator groups:")
for n, r in nodes.nlargest(5, "n_predators").iterrows():
    log(f"  {n:28s} {r.n_predators}")
log("largest summed diet share (out-strength):")
for n, r in nodes.nlargest(5, "diet_share_sum").iterrows():
    log(f"  {n:28s} {r.diet_share_sum:.2f}")
log("highest trophic level:")
for n, r in nodes.nlargest(5, "trophic_level").iterrows():
    log(f"  {n:28s} {r.trophic_level:.2f}")
log(f"\ngroups in feeding loops (strongly connected, size > 1): {len(in_loop)}")

# How fragile is the web to the modelling choice of keeping detritus?
H = G.copy()
H.remove_nodes_from(NONLIVING)
log(f"\nwithout Detritus/Discard: {H.number_of_nodes()} nodes, "
    f"{H.number_of_edges()} edges, isolates {list(nx.isolates(H)) or 'none'}, "
    f"weak components {nx.number_weakly_connected_components(H)}")

# ---------------------------------------------------------------- save
culled = pd.DataFrame(M, index=pd.Index(names, name="prey \\ predator"),
                      columns=names)
edges = pd.DataFrame(
    [(u, v, d["weight"], D[idx[u], idx[v]]) for u, v, d in G.edges(data=True)],
    columns=["prey", "predator", "diet_share", "diet_share_renorm"])
edges["diet_share_renorm"] = edges["diet_share_renorm"].round(6)
edges = edges.sort_values(["predator", "diet_share"], ascending=[True, False])
log_df = pd.DataFrame(cleaning)

culled.to_csv(OUT / "cadiz_diet_matrix_culled.csv")
log_df.to_csv(OUT / "cadiz_cleaning_log.csv", index=False)
nodes.to_csv(OUT / "cadiz_node_traits.csv")
edges.to_csv(OUT / "cadiz_edges.csv", index=False)
for n in G:
    G.nodes[n]["trophic_level"] = float(nodes.at[n, "trophic_level"])
nx.write_graphml(G, OUT / "cadiz_foodweb.graphml")
with open(OUT / "summary.txt", "w") as f:
    f.write("\n".join(report) + "\n")

# Everything above as one workbook, for anyone who would rather use Excel.
COLUMNS = [
    ("group_name", "Functional group (node). Groups, not species: some are one species, some are pools."),
    ("category", "Broad type, assigned by us from the group name."),
    ("living", "False for Detritus and Discard."),
    ("role", "basal = eats nothing; top predator = eaten by no other group; otherwise intermediate. Cannibalism ignored."),
    ("trophic_level", "Ecopath-style: 1 + diet-weighted mean trophic level of the prey; basal groups = 1. Uses renormalised diets."),
    ("omnivory_index", "Diet-weighted variance of prey trophic levels (Christensen & Pauly 1992). 0 = feeds on a single level."),
    ("n_prey", "Number of prey groups (in-degree, generality), not counting itself."),
    ("n_predators", "Number of predator groups (out-degree, vulnerability), not counting itself."),
    ("total_degree", "n_prey + n_predators."),
    ("effective_prey", "Diet diversity: exp(Shannon entropy) of the renormalised diet without cannibalism. Blank for basal groups."),
    ("diet_share_sum", "Sum of this group's share across all predator diets (out-strength): importance as prey."),
    ("cannibal", "True if the group eats itself (self-loop)."),
    ("cannibalism_share", "Share of the group's own diet that is itself."),
    ("import_share_removed", "Share of the diet that was 'Import' (eaten outside the area) and removed in cleaning."),
    ("diet_total_original", "Sum of the group's diet column before cleaning (1 for consumers, 0 for basal)."),
    ("in_feeding_loop", "True if the group sits in a feeding loop (strongly connected component larger than one)."),
]
SHEETS = [
    ("culled_diet_matrix", "Cleaned 43 x 43 diet matrix. Rows = prey, columns = predators, cell = share of the predator's diet. Original shares; Import removed."),
    ("cleaning_log", "Every change made to the original download, in order, with link counts before and after."),
    ("node_traits", "One row per group: trophic level and other derived traits. Column definitions below."),
    ("edge_list", "One row per feeding link (prey -> predator), with original and renormalised diet share."),
    ("original_matrix", "The 44 x 44 matrix exactly as downloaded from Web of Life on 2026-10-15."),
]
readme = [["Gulf of Cadiz food web: cleaned data and derived traits"],
          ["Source: Torres et al. 2013, Ecological Modelling 265:26-44, doi:10.1016/j.ecolmodel.2013.05.019, via the Web of Life database (www.web-of-life.es), network FW_005."],
          ["Generated by code/build_network.py. Do not edit by hand: change the code and rerun it, or the workbook and the CSVs in output/ will disagree."],
          [], ["Sheet", "Contents"], *[[s, d] for s, d in SHEETS],
          [], ["node_traits column", "Definition"], *[[c, d] for c, d in COLUMNS]]

xlsx = OUT / "cadiz_foodweb_data.xlsx"
with pd.ExcelWriter(xlsx, engine="openpyxl") as xw:
    pd.DataFrame(readme).to_excel(xw, sheet_name="README", index=False, header=False)
    culled.to_excel(xw, sheet_name="culled_diet_matrix")
    log_df.to_excel(xw, sheet_name="cleaning_log", index=False)
    nodes.reset_index().to_excel(xw, sheet_name="node_traits", index=False)
    edges.to_excel(xw, sheet_name="edge_list", index=False)
    raw.rename_axis("prey \\ predator").to_excel(xw, sheet_name="original_matrix")

from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter
wb = load_workbook(xlsx)
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(name="Arial", size=10, bold=c.font.bold)
ws = wb["README"]
ws["A1"].font = Font(name="Arial", size=12, bold=True)
for r in ws.iter_rows():
    if r[0].value in ("Sheet", "node_traits column"):
        for c in r:
            c.font = Font(name="Arial", size=10, bold=True)
ws.column_dimensions["A"].width = 24
ws.column_dimensions["B"].width = 110
for name in ("cleaning_log", "node_traits", "edge_list"):
    ws = wb[name]
    ws.freeze_panes = "B2"
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(width + 2, 60)
for c in wb["cleaning_log"]["D"]:
    c.alignment = Alignment(wrap_text=True, vertical="top")
wb["cleaning_log"].column_dimensions["D"].width = 90
for name in ("culled_diet_matrix", "original_matrix"):
    ws = wb[name]
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 28
    for k in range(2, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(k)].width = 9
        ws.cell(1, k).alignment = Alignment(text_rotation=90)
    rng = f"B2:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.conditional_formatting.add(rng, ColorScaleRule(
        start_type="num", start_value=0, start_color="FFFFFF",
        end_type="num", end_value=1, end_color="1C5CAB"))
ws = wb["original_matrix"]                  # shade the row that was removed
for c in ws[raw.index.get_loc("Import") + 2]:
    c.fill = PatternFill("solid", fgColor="F4CCCC")
wb.save(xlsx)

# ---------------------------------------------------------------- figures
INK, INK2, GRID = "#0b0b0b", "#52514e", "#e1e0d9"
GROUP_COLOR = {"invertebrates & plankton": "#2a78d6", "fishes": "#eb6834",
               "marine mammals, birds, turtles": "#1baf7a",
               "non-living": "#898781"}
plt.rcParams.update({"font.family": "sans-serif", "font.size": 9,
                     "text.color": INK, "axes.labelcolor": INK2,
                     "xtick.color": INK2, "ytick.color": INK2})
group_handles = [Line2D([], [], marker="s", ls="", color=c, markersize=8, label=g)
                 for g, c in GROUP_COLOR.items()]

# 1. Diet matrix, both axes ordered by trophic level. Group identity sits in
# a colour strip beside each axis so the labels can stay in ink.
order = list(nodes.sort_values("trophic_level").index)
A = M[np.ix_([idx[n] for n in order], [idx[n] for n in order])]
cmap = LinearSegmentedColormap.from_list(
    "blue", ["#cde2fb", "#86b6ef", "#3987e5", "#1c5cab", "#0d366b"])
strip = np.array([[to_rgb(GROUP_COLOR[group(n)])] for n in order])

fig, ax = plt.subplots(figsize=(10.5, 10))
im = ax.imshow(np.ma.masked_equal(A, 0), cmap=cmap,
               norm=LogNorm(vmin=1e-4, vmax=1), interpolation="none")
div = make_axes_locatable(ax)
sl = div.append_axes("left", size="1.6%", pad=0.04, sharey=ax)
sb = div.append_axes("bottom", size="1.6%", pad=0.04, sharex=ax)
sl.imshow(strip, aspect="auto", interpolation="none")
sb.imshow(strip.transpose(1, 0, 2), aspect="auto", interpolation="none")
sl.set_yticks(range(N), order)
sl.set_xticks([])
sb.set_xticks(range(N), order, rotation=90)
sb.set_yticks([])
sl.set_ylabel("prey (resource)")
sb.set_xlabel("predator (consumer)")
ax.tick_params(labelleft=False, labelbottom=False, left=False, bottom=False)
for a in (sl, sb):
    a.tick_params(length=0)
    for sp in a.spines.values():
        sp.set_visible(False)
ax.set_xticks(np.arange(-.5, N), minor=True)
ax.set_yticks(np.arange(-.5, N), minor=True)
ax.grid(which="minor", color=GRID, lw=0.4)
ax.tick_params(which="minor", length=0)
for sp in ax.spines.values():
    sp.set_color(GRID)
for k, n in enumerate(order):               # outline cannibalism cells
    if n in loops:
        ax.add_patch(Rectangle((k - .5, k - .5), 1, 1, fill=False,
                               edgecolor=INK, lw=0.9))
cax = div.append_axes("right", size="2.2%", pad=0.2)
cb = fig.colorbar(im, cax=cax)
cb.set_label("share of predator's diet (log scale; < 1e-4 drawn at floor)")
cb.outline.set_visible(False)
handles = group_handles + [Line2D([], [], marker="s", ls="", markersize=9,
                                  markerfacecolor="none", markeredgecolor=INK,
                                  label="cannibalism (diagonal)")]
fig.legend(handles=handles, loc="lower center", ncol=5, frameon=False,
           fontsize=8.5, handletextpad=0.3)
fig.subplots_adjust(left=0.24, right=0.9, top=0.98, bottom=0.25)
fig.savefig(FIG / "cadiz_diet_matrix.png", dpi=200)
plt.close(fig)

# 2. Food web: y = trophic level, x from a 1-D force layout (y held fixed),
# then nodes are pushed sideways until no two labels overlap.
rng = np.random.default_rng(4)
y = {n: TL[idx[n]] for n in G}
x = {n: rng.uniform(-1, 1) for n in G}
U = G.to_undirected()
for _ in range(400):
    for n in G:
        pull = np.mean([x[m] for m in U[n] if m != n])
        push = sum(np.sign(x[n] - x[m]) * 0.012 / (abs(x[n] - x[m]) + 0.05)
                   for m in G if m != n and abs(y[n] - y[m]) < 0.35)
        x[n] = 0.85 * x[n] + 0.15 * pull + push
xs = np.array(list(x.values()))
x = {n: (x[n] - xs.min()) / (xs.max() - xs.min()) for n in G}

strength = pd.Series({n: G.out_degree(n, weight="weight") for n in names})
size = 40 + 260 * np.sqrt(strength / strength.max())
fig, ax = plt.subplots(figsize=(12, 9.5))
fig.subplots_adjust(left=0.06, right=0.98, top=0.98, bottom=0.1)
ax.set_xlim(-0.35, 1.35)
ax.set_ylim(0.8, TL.max() + 0.25)
labels = {n: ax.annotate(n, (x[n], y[n]), xytext=(0, 10 if n in loops else 7),
                         textcoords="offset points", ha="center",
                         va="bottom", fontsize=7.5, color=INK, zorder=4)
          for n in G}
px_per_x = ax.transData.transform((1, 0))[0] - ax.transData.transform((0, 0))[0]
def node_box(n):
    r = (np.sqrt(size[n] + (110 if n in loops else 0)) / 2 + 2) * fig.dpi / 72
    cx, cy = ax.transData.transform((x[n], y[n]))
    return Bbox.from_extents(cx - r, cy - r, cx + r, cy + r)

for _ in range(300):
    fig.canvas.draw()
    # each node owns its marker and its label; no box may touch another node's
    box = {n: [labels[n].get_window_extent().expanded(1.06, 1.25), node_box(n)]
           for n in G}
    moved = False
    for a, b in combinations(G, 2):
        for ba in box[a]:
            for bb in box[b]:
                if ba.overlaps(bb):
                    gap = min(ba.x1, bb.x1) - max(ba.x0, bb.x0)
                    d = (gap / 2 + 1) / px_per_x
                    lo, hi = (a, b) if (x[a], a) < (x[b], b) else (b, a)
                    x[lo] -= d
                    x[hi] += d
                    moved = True
                    break
            else:
                continue
            break
    for n in G:
        labels[n].xy = (x[n], y[n])
    if not moved:
        break
assert not moved, "labels still overlap"
pos = {n: (x[n], y[n]) for n in G}

for u, v, d in sorted(G.edges(data=True), key=lambda e: e[2]["weight"]):
    if u == v:
        continue
    wt = d["weight"]
    ax.annotate("", xy=pos[v], xytext=pos[u],
                arrowprops=dict(arrowstyle="-|>", mutation_scale=6 + 6 * wt,
                                lw=0.25 + 3.2 * wt, color=INK2,
                                alpha=0.07 + 0.6 * wt, shrinkA=6, shrinkB=6,
                                connectionstyle="arc3,rad=0.08"))
for n in G:
    ax.scatter(*pos[n], s=size[n], color=GROUP_COLOR[group(n)],
               edgecolor="white", linewidth=1.2, zorder=3)
    if n in loops:                          # ring marks cannibalism
        ax.scatter(*pos[n], s=size[n] + 110, facecolor="none",
                   edgecolor=INK, linewidth=0.9, zorder=2)
ax.set_ylabel("trophic level")
ax.set_xticks([])
for s in ("top", "right", "bottom"):
    ax.spines[s].set_visible(False)
ax.spines["left"].set_color(GRID)
ax.grid(axis="y", color=GRID, lw=0.5)
ax.set_axisbelow(True)
xs = np.array([p[0] for p in pos.values()])
ax.set_xlim(xs.min() - 0.08, xs.max() + 0.08)   # zooming in only adds space
handles = [Line2D([], [], marker="o", ls="", color=c, markersize=8, label=g)
           for g, c in GROUP_COLOR.items()]
handles += [Line2D([], [], marker="o", ls="", markerfacecolor="none",
                   markeredgecolor=INK, markersize=10,
                   label="cannibalism (self-loop)"),
            Line2D([], [], color=INK2, lw=2.5, alpha=0.6,
                   label="prey → predator, width = diet share")]
fig.legend(handles=handles, loc="lower center", ncol=3, frameon=False,
           fontsize=8.5)
fig.savefig(FIG / "cadiz_foodweb.png", dpi=200)
plt.close(fig)
