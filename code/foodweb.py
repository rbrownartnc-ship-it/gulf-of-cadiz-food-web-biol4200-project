"""Shared helpers for the food-web analyses.

Imported by distributions.py, network_structure.py, node_removal.py,
edge_removal.py and the checks in code/checks/. Everything here works on the
cleaned web that build_network.py writes to output/, so run that first.
"""
from pathlib import Path
import numpy as np
import pandas as pd
import networkx as nx
import matplotlib.pyplot as plt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"
FIG = ROOT / "figures"
NONLIVING = {"Detritus", "Discard"}
SEED = 4200

# ---------------------------------------------------------------- data
def load_web():
    """The cleaned Gulf of Cadiz web as a DiGraph with edges prey -> predator.

    Edge attributes: w = original diet share, d = diet share rescaled so each
    predator's diet sums to 1. Node attributes: category, trophic_level, living.
    Returns (G, traits) where traits is output/cadiz_node_traits.csv.
    """
    edges = pd.read_csv(OUT / "cadiz_edges.csv")
    traits = pd.read_csv(OUT / "cadiz_node_traits.csv", index_col=0)
    G = nx.DiGraph()
    for n, r in traits.iterrows():
        G.add_node(n, category=r.category, trophic_level=r.trophic_level,
                   living=bool(r.living))
    for r in edges.itertuples():
        G.add_edge(r.prey, r.predator, w=r.diet_share, d=r.diet_share_renorm)
    assert G.number_of_nodes() == 43 and G.number_of_edges() == 411
    return G, traits


def simple(G):
    """The same web without cannibalistic self-loops and without weights."""
    H = nx.DiGraph()
    H.add_nodes_from(G)
    H.add_edges_from((u, v) for u, v in G.edges if u != v)
    return H

# ---------------------------------------------------------------- extinctions
class Cascade:
    """Secondary extinctions after removing groups from a food web.

    A consumer goes secondarily extinct once it has lost at least `theta` of
    its diet (rescaled diet shares; cannibalism never keeps a group alive).
      theta = 1   topological rule (Dunne, Williams & Martinez 2002): a
                  consumer dies only when every one of its prey is gone.
      theta < 1   weighted rule: e.g. 0.5 = dies after losing half its diet.
    Groups with no prey (phytoplankton, detritus, discard) never starve; they
    disappear only when removed directly.
    """

    def __init__(self, G):
        self.G = G
        self.diet = {n: {p: G[p][n]["d"] for p in G.predecessors(n) if p != n}
                     for n in G}

    def run(self, removed, theta=1.0, present=None, removed_links=()):
        """Remove the groups `removed` and the feeding links `removed_links`
        (pairs (prey, predator)) from `present` (default: every group), and
        return (still_present, secondary_extinctions)."""
        present = set(self.G if present is None else present) - set(removed)
        cut = set(removed_links)
        secondary = set()
        changed = True
        while changed:
            changed = False
            for n in sorted(present):
                diet = self.diet[n]
                if not diet:                         # basal: never starves
                    continue
                left = sum(s for p, s in diet.items() if p in present and (p, n) not in cut)
                if left <= (1 - theta) + 1e-12:
                    present.discard(n)
                    secondary.add(n)
                    changed = True
        return present, secondary


class LinkRemoval:
    """Incremental version of Cascade for cutting feeding links one at a time
    (hundreds of steps), with the same extinction rule: a consumer starves
    once its uncut links from living prey carry no more than 1 - theta of its
    diet. Losing predators never kills a group, so basal groups always stay.
    theta = None cuts links without any starvation (pure link percolation)."""

    def __init__(self, cascade, theta):
        self.diet, self.theta = cascade.diet, theta
        self.alive = set(cascade.G)
        self.links = {(p, n) for n, d in self.diet.items() for p in d}
        self.predators = {n: [] for n in cascade.G}
        for n, d in self.diet.items():
            for p in d:
                self.predators[p].append(n)
        self.secondary = set()

    def _starving(self, n):
        diet = self.diet[n]
        if not diet or self.theta is None:       # theta None: cut links, nothing starves
            return False
        left = sum(s for p, s in diet.items() if p in self.alive and (p, n) in self.links)
        return left <= (1 - self.theta) + 1e-12

    def _die(self, n):
        stack = [n]
        while stack:
            m = stack.pop()
            if m not in self.alive:
                continue
            self.alive.discard(m)
            self.secondary.add(m)
            for q in self.predators[m]:
                if q in self.alive and (m, q) in self.links and self._starving(q):
                    stack.append(q)

    def remove_link(self, prey, predator):
        """Cut one link and let starvation cascade."""
        if (prey, predator) not in self.links:
            return
        self.links.discard((prey, predator))
        if predator in self.alive and prey in self.alive and self._starving(predator):
            self._die(predator)

    def graph(self, G):
        """The surviving web: living groups and uncut links."""
        H = nx.DiGraph()
        H.add_nodes_from(self.alive)
        H.add_edges_from((p, n) for p, n in self.links if p in self.alive and n in self.alive)
        return H


def largest_wcc(G, present):
    """Size of the largest weakly connected component among `present`."""
    H = G.subgraph(present)
    return max((len(c) for c in nx.weakly_connected_components(H)), default=0)


def largest_scc(G, present):
    """Size of the largest strongly connected component (a feeding loop) among
    `present`; a lone group counts as 1."""
    H = G.subgraph(present)
    return max((len(c) for c in nx.strongly_connected_components(H)), default=0)

# ---------------------------------------------------------------- binary structure
def prey_averaged_tl(H):
    """Unweighted trophic level: 1 + mean trophic level of the prey (self-loops
    ignored), basal = 1. Used for the null-model comparison, where webs have no
    weights. Returns None if some group cannot be reached from a basal group."""
    nodes = list(H)
    k = {n: sum(1 for p in H.predecessors(n) if p != n) for n in nodes}
    i = {n: j for j, n in enumerate(nodes)}
    A = np.eye(len(nodes))
    for u, v in H.edges:
        if u != v:
            A[i[v], i[u]] -= 1 / k[v]
    try:
        tl = np.linalg.solve(A, np.ones(len(nodes)))
    except np.linalg.LinAlgError:
        return None
    if not np.all(np.isfinite(tl)) or tl.max() > 100:
        return None
    return dict(zip(nodes, tl))


def shortest_path_tl(H):
    """1 + length of the shortest food chain down to any basal group."""
    basal = [n for n in H if not any(p != n for p in H.predecessors(n))]
    dist = nx.multi_source_dijkstra_path_length(H, basal) if basal else {}
    return {n: 1 + dist[n] for n in dist}


def food_web_properties(G):
    """Standard food-web descriptors (Williams & Martinez 2000) for a web with
    edges prey -> predator. Self-loops count as cannibalism."""
    S = G.number_of_nodes()
    L = G.number_of_edges()
    H = simple(G)
    n_prey = {n: sum(1 for p in G.predecessors(n) if p != n) for n in G}
    n_pred = {n: sum(1 for q in G.successors(n) if q != n) for n in G}
    basal = [n for n in G if n_prey[n] == 0]
    top = [n for n in G if n_pred[n] == 0 and n_prey[n] > 0]
    sptl = shortest_path_tl(H)
    omn = [n for n in G if n_prey[n] > 1 and
           len({sptl.get(p) for p in G.predecessors(n) if p != n}) > 1]
    reachable = len(sptl) == S          # every group sits on a chain from a basal group
    tl = prey_averaged_tl(H)
    in_loop = sum(len(c) for c in nx.strongly_connected_components(H) if len(c) > 1)
    U = H.to_undirected()
    giant = U.subgraph(max(nx.connected_components(U), key=len))
    gen = np.array([G.in_degree(n) for n in G]) / (L / S)
    vul = np.array([G.out_degree(n) for n in G]) / (L / S)
    mutual = sum(1 for u, v in H.edges if H.has_edge(v, u)) / H.number_of_edges()
    # Trophic coherence (Johnson et al. 2014): SD over links of the difference
    # in (unweighted) trophic level between predator and prey.
    q = np.std([tl[v] - tl[u] for u, v in H.edges]) if tl else np.nan
    return {
        "connectance": L / S**2,
        "link_density": L / S,
        "frac_basal": len(basal) / S,
        "frac_intermediate": (S - len(basal) - len(top)) / S,
        "frac_top": len(top) / S,
        "frac_cannibal": sum(1 for n in G if G.has_edge(n, n)) / S,
        "frac_omnivore": len(omn) / S if reachable else np.nan,
        "frac_in_loops": in_loop / S,
        "mean_trophic_level": np.mean(list(tl.values())) if tl else np.nan,
        "max_trophic_level": max(tl.values()) if tl else np.nan,
        "generality_sd": gen.std(),
        "vulnerability_sd": vul.std(),
        "clustering": nx.average_clustering(U),
        "path_length": nx.average_shortest_path_length(giant),
        "frac_mutual_links": mutual,
        "assortativity": nx.degree_assortativity_coefficient(U),
        "trophic_coherence_q": q,
    }


def niche_model(S, C, rng):
    """One niche-model web (Williams & Martinez 2000) with S species and
    expected directed connectance C, edges prey -> predator."""
    beta = 1 / (2 * C) - 1
    n = rng.uniform(0, 1, S)
    r = n * rng.beta(1, beta, S)
    r[np.argmin(n)] = 0                       # lowest niche value is basal
    c = rng.uniform(r / 2, n)
    lo, hi = c - r / 2, c + r / 2
    eats = (n[None, :] >= lo[:, None]) & (n[None, :] <= hi[:, None])   # [i, j]: i eats j
    G = nx.DiGraph()
    G.add_nodes_from(range(S))
    G.add_edges_from((j, i) for i, j in zip(*np.nonzero(eats)))
    return G


def acceptable_niche_web(G, C_target, tol=0.03):
    """The usual acceptance rules: connectance within tol of the target, one
    connected web, and every species reachable from a basal species."""
    S, L = G.number_of_nodes(), G.number_of_edges()
    if abs(L / S**2 - C_target) > tol * C_target:
        return False
    if not nx.is_weakly_connected(G):
        return False
    return prey_averaged_tl(simple(G)) is not None and len(shortest_path_tl(simple(G))) == S

# ---------------------------------------------------------------- importance
def pagerank_importance(H):
    """Allesina & Pascual (2009): PageRank on the web with links pointing from
    consumer to resource, plus a 'root' that every basal group points to and
    that points to every group. A group is important if important groups
    depend on it. The root is dropped and the scores rescaled to sum to 1."""
    R = H.reverse(copy=True)
    R.add_node("_root")
    for n in H:
        R.add_edge("_root", n)
        if not any(p != n for p in H.predecessors(n)):
            R.add_edge(n, "_root")
    pr = nx.pagerank(R, alpha=0.99, max_iter=10_000, tol=1e-12)
    pr.pop("_root")
    s = sum(pr.values())
    return {n: v / s for n, v in pr.items()}


# ---------------------------------------------------------------- figures
# Same conventions as build_network.py: drawn at printed size (7 in = full text
# width of the two-column template, 3.4 in = one column), text >= 6.5 pt,
# saved as vector PDF + 300-dpi PNG. One idea per figure; figures live in
# numbered subfolders of figures/ by analysis.
COLUMN, FULL = 3.4, 7.0
INK, INK2, GRID = "#0b0b0b", "#52514e", "#e1e0d9"
BLUE, ORANGE, AQUA, GREY = "#2a78d6", "#eb6834", "#1baf7a", "#898781"
CATEGORY_STYLE = {   # color + marker, so category never rests on color alone
    "invertebrates & plankton": (BLUE, "o"),
    "fishes": (ORANGE, "s"),
    "marine mammals, birds, turtles": (AQUA, "^"),
    "non-living": (GREY, "X"),
}

def use_style():
    plt.rcParams.update({"font.family": "sans-serif", "font.size": 7.5,
                         "axes.labelsize": 7.5, "axes.titlesize": 7.5,
                         "xtick.labelsize": 6.5, "ytick.labelsize": 6.5,
                         "legend.fontsize": 6.5, "text.color": INK,
                         "axes.labelcolor": INK2, "xtick.color": INK2,
                         "ytick.color": INK2, "axes.edgecolor": GRID,
                         "axes.spines.top": False, "axes.spines.right": False,
                         "pdf.fonttype": 42})


def save(fig, folder, stem, width=COLUMN):
    """Save figures/<folder>/<stem>.pdf and .png, then close the figure.
    The canvas is narrowed until the cropped figure (labels included) is
    exactly `width` inches, so it prints at that width without shrinking."""
    for _ in range(6):
        w = fig.get_tightbbox(fig.canvas.get_renderer()).width + 0.08   # + pad_inches on both sides
        if abs(w - width) < 0.01:
            break
        fig.set_figwidth(fig.get_figwidth() - (w - width))
    d = FIG / folder
    d.mkdir(parents=True, exist_ok=True)
    fig.savefig(d / f"{stem}.pdf", bbox_inches="tight", pad_inches=0.04,
                metadata={"CreationDate": None})
    fig.savefig(d / f"{stem}.png", dpi=300, bbox_inches="tight", pad_inches=0.04)
    plt.close(fig)


def plain_log_ticks(axis):
    """Tick labels on a log axis as plain numbers (0.01, 0.1, 1; 1e−6 for tiny
    values) instead of 10^n, whose superscripts print below 5 pt."""
    from matplotlib.ticker import FuncFormatter, LogLocator, NullFormatter
    def fmt(v, _):
        if v <= 0:
            return ""
        return f"{v:g}" if v >= 0.001 else f"1e−{round(-np.log10(v))}"
    axis.set_major_locator(LogLocator(base=10))
    axis.set_major_formatter(FuncFormatter(fmt))
    axis.set_minor_formatter(NullFormatter())


def wrap_label(text, width=30):
    """Break a long axis label onto at most two lines at a space."""
    import textwrap
    lines = textwrap.wrap(text, width)
    return text if len(lines) <= 1 else lines[0] + "\n" + " ".join(lines[1:])


def tidy(ax, grid="both"):
    """Recessive grid and ticks, shared by every single-panel figure."""
    if grid:
        ax.grid(axis=grid, color=GRID, lw=0.4)
    ax.set_axisbelow(True)
    ax.tick_params(length=2, pad=1.5)


def category_legend(fig, y=-0.02, ncol=2):
    """Legend for color + marker by broad category, below the axes."""
    from matplotlib.lines import Line2D
    h = [Line2D([], [], ls="", marker=mk, color=c, markersize=4.5, label=cat)
         for cat, (c, mk) in CATEGORY_STYLE.items()]
    fig.legend(handles=h, loc="upper center", bbox_to_anchor=(0.5, y), ncol=ncol,
               frameon=False, handletextpad=0.2, columnspacing=1.0)

# ---------------------------------------------------------------- workbook
RESULT_SHEETS = [   # (sheet, csv in output/, description)
    ("distribution_stats", "cadiz_distribution_stats.csv",
     "Summary statistics of every node and link distribution (mean, median, SD, skew, dispersion)."),
    ("global_structure", "cadiz_global_structure.csv",
     "Whole-web measures: size, connectance, components, clustering, path length."),
    ("structure_vs_null", "cadiz_structure_vs_null.csv",
     "Empirical web vs random (ER), degree-preserving and niche-model webs: null mean, SD, z-score, percentile."),
    ("centrality", "cadiz_centrality.csv",
     "One row per group: degree, strength, betweenness, clustering, PageRank importance."),
    ("centrality_vs_tl", "cadiz_centrality_vs_tl.csv",
     "Spearman correlation of each centrality with trophic level (consumers, and all groups)."),
    ("node_knockouts", "cadiz_node_knockouts.csv",
     "Remove one group at a time: secondary extinctions and component sizes."),
    ("guild_removal", "cadiz_guild_removal.csv",
     "Remove a whole guild (trophic level band or taxonomic group) at once."),
    ("node_robustness_summary", "cadiz_node_robustness_summary.csv",
     "R50 for each group-removal order and extinction rule."),
    ("node_robustness_curves", "cadiz_node_robustness_curves.csv",
     "Step-by-step group-removal curves behind the robustness figures."),
    ("node_fragmentation", "cadiz_node_fragmentation.csv",
     "When the surviving web first splits into pieces during group removal."),
    ("link_knockouts", "cadiz_link_knockouts.csv",
     "Cut one feeding link at a time: secondary extinctions, edge betweenness, bridges."),
    ("link_robustness_summary", "cadiz_link_robustness_summary.csv",
     "L50 for each link-removal order and extinction rule."),
    ("link_robustness_curves", "cadiz_link_robustness_curves.csv",
     "Step-by-step link-removal curves behind the link figures."),
    ("link_fragmentation", "cadiz_link_fragmentation.csv",
     "When the surviving web first splits into pieces while links are cut."),
]

def update_results_workbook():
    """(Re)build output/cadiz_analysis_results.xlsx from whichever result CSVs
    exist, one sheet each, so either analysis script can run first."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font, Alignment
    present = [(s, f, d) for s, f, d in RESULT_SHEETS if (OUT / f).exists()]
    path = OUT / "cadiz_analysis_results.xlsx"
    readme = [["Gulf of Cadiz food web: network analysis results"],
              ["Generated by the analysis scripts in code/ (see code/README.md). "
               "Do not edit by hand; rerun the code."], [],
              ["Sheet", "Contents", "Same data as"]]
    readme += [[s, d, f"output/{f}"] for s, f, d in present]
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame(readme).to_excel(xw, sheet_name="README", index=False, header=False)
        for s, f, _ in present:
            pd.read_csv(OUT / f).to_excel(xw, sheet_name=s, index=False)
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                c.font = Font(name="Arial", size=10, bold=c.font.bold)
        if ws.title != "README":
            ws.freeze_panes = "B2"
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(width + 2, 70)
    ws = wb["README"]
    ws["A1"].font = Font(name="Arial", size=12, bold=True)
    for c in ws[4]:
        c.font = Font(name="Arial", size=10, bold=True)
    for row in ws.iter_rows(min_row=5):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["B"].width = 90
    wb.save(path)
